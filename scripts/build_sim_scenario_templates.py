from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from backend.app.services.excel_io import (  # noqa: E402
    EQUIPMENT_GROUP_COLUMNS,
    SPECTRUM_RULE_COLUMNS,
    TASK_UNIT_COLUMNS,
)
from backend.app.services.sim_scenarios import SIM_SCENARIO_MANIFESTS, sim_scenario_records  # noqa: E402


OUTPUT_ROOT = ROOT / "samples" / "sim_scenarios"


def _write_xlsx(path: Path, rows: list[dict], columns: list[str], table_name: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.freeze_panes = "A2"
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(column) for column in columns])

    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, column in enumerate(columns, start=1):
        values = [str(sheet.cell(row, index).value or "") for row in range(1, min(sheet.max_row, 80) + 1)]
        sheet.column_dimensions[get_column_letter(index)].width = min(42, max(12, max(map(len, values), default=10) + 2))

    if sheet.max_row >= 2:
        table = Table(displayName=table_name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    sheet.sheet_view.showGridLines = False
    workbook.save(path)


def build_templates(output_root: Path = OUTPUT_ROOT) -> list[Path]:
    generated: list[Path] = []
    output_root.mkdir(parents=True, exist_ok=True)
    for scenario_key, metadata in SIM_SCENARIO_MANIFESTS.items():
        scenario_dir = output_root / scenario_key
        scenario_dir.mkdir(parents=True, exist_ok=True)
        task_units, equipment_groups, spectrum_rules = sim_scenario_records(scenario_key)

        task_path = scenario_dir / "task_units.xlsx"
        equipment_path = scenario_dir / "equipment_groups.xlsx"
        rules_path = scenario_dir / "spectrum_rules.xlsx"
        _write_xlsx(task_path, task_units, TASK_UNIT_COLUMNS, f"TaskUnits{len(generated):02d}")
        _write_xlsx(equipment_path, equipment_groups, EQUIPMENT_GROUP_COLUMNS, f"EquipmentGroups{len(generated):02d}")
        _write_xlsx(rules_path, spectrum_rules, SPECTRUM_RULE_COLUMNS, f"SpectrumRules{len(generated):02d}")

        manifest = {
            **metadata,
            "scenario_key": scenario_key,
            "data_boundary": "Public sources define mission structure only; all operational inputs are synthetic.",
            "files": {
                "task_units": task_path.name,
                "equipment_groups": equipment_path.name,
                "spectrum_rules": rules_path.name,
            },
            "counts": {
                "task_units": len(task_units),
                "equipment_groups": len(equipment_groups),
                "spectrum_rules": len(spectrum_rules),
                "equipment_samples": sum(int(row.get("count") or 0) for row in equipment_groups),
            },
        }
        manifest_path = scenario_dir / "manifest.json"
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        generated.extend([task_path, equipment_path, rules_path, manifest_path])
    return generated


if __name__ == "__main__":
    paths = build_templates()
    print(f"generated={len(paths)} root={OUTPUT_ROOT.relative_to(ROOT)}")
