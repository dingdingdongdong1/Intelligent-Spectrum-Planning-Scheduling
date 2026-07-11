from __future__ import annotations

import json
from io import BytesIO

from openpyxl import load_workbook
from sqlmodel import Session, SQLModel, create_engine, select

from backend.app.models import AuditLog, EquipmentAssignment, Project
from backend.app.services.equipment_planning import (
    TASK_OBJECTIVES,
    TASK_SCENARIOS,
    _available_ranges_from_message,
    _forbidden_ranges_from_message,
    _agent_next_actions,
    _agent_strategy_history,
    _parse_resource_segments,
    _replan_change_items,
    _reuse_risks,
    _sanitize_trial_context,
    _strategy_constraint_variants,
    _supplement_range_for_unmet_assignment,
    _supplement_ranges_for_unmet_assignments,
    _trial_prediction_alignment,
    _weighted_plan_decision,
    build_task_agent_assessment,
    build_task_strategy_trials,
    build_task_visualization_data,
    demo_scenario_records,
    execute_capacity_batch_planning,
    execute_capacity_batch_risk_closure,
    latest_capacity_batch_closure_export_xlsx,
    latest_capacity_batch_export_xlsx,
    parametric_demo_records,
    replace_equipment_groups,
    replace_spectrum_rules,
    replace_task_units,
    run_task_planning,
    solve_task_assignment,
    task_planning_batch_performance_test,
    task_planning_performance_test,
    task_versions_and_audit,
    task_sample_catalog,
    validate_task_inputs,
)


def test_demo_scenario_has_expected_scale() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    assert len(task_units) == 12
    assert len(equipment_groups) == 20
    assert sum(item["count"] for item in equipment_groups) >= 300
    assert any(item["rule_type"] == "禁用" for item in spectrum_rules)
    assert any(item["rule_type"] == "保护" for item in spectrum_rules)


def test_task_validation_accepts_demo_scenario() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    result = validate_task_inputs(task_units, equipment_groups, spectrum_rules)
    assert result["ok"] is True
    assert result["summary"]["equipment_group_count"] == 20


def test_task_solver_outputs_partial_or_full_decisions() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    result = solve_task_assignment(task_units, equipment_groups, spectrum_rules, "task_assurance")
    assert result["status"] == "success"
    assert len(result["assignments"]) == len(equipment_groups)
    assert {item["status"] for item in result["assignments"]} <= {"完全满足", "部分满足", "未满足"}
    assert result["summary"]["task_satisfaction_avg"] > 0


def test_task_conflict_model_covers_same_frequency_spacing_power_and_intermodulation() -> None:
    task_units = [
        {"task_unit_id": "TU-A", "area_center_lat": 30.0, "area_center_lon": 110.0, "priority": 5, "spectrum_relation": "独占"},
        {"task_unit_id": "TU-B", "area_center_lat": 30.01, "area_center_lon": 110.01, "priority": 4, "spectrum_relation": "可复用"},
        {"task_unit_id": "TU-C", "area_center_lat": 30.02, "area_center_lon": 110.02, "priority": 3, "spectrum_relation": "可复用"},
    ]
    equipment_groups = [
        {"equipment_group_id": "EG-A", "task_unit_id": "TU-A", "tx_power_w": 150, "protection_distance_km": 20, "min_spacing_khz": 50, "guard_band_khz": 50, "bandwidth_khz": 25, "priority": 5},
        {"equipment_group_id": "EG-B", "task_unit_id": "TU-B", "tx_power_w": 120, "protection_distance_km": 20, "min_spacing_khz": 50, "guard_band_khz": 50, "bandwidth_khz": 25, "priority": 4},
        {"equipment_group_id": "EG-C", "task_unit_id": "TU-C", "tx_power_w": 20, "protection_distance_km": 10, "min_spacing_khz": 50, "guard_band_khz": 50, "bandwidth_khz": 25, "priority": 3},
        {"equipment_group_id": "EG-D", "task_unit_id": "TU-B", "tx_power_w": 20, "protection_distance_km": 10, "min_spacing_khz": 50, "guard_band_khz": 50, "bandwidth_khz": 25, "priority": 3},
    ]
    assignments = [
        {"task_unit_id": "TU-A", "equipment_group_id": "EG-A", "band_group": "TEST", "assigned_resource": "100.000000 MHz"},
        {"task_unit_id": "TU-B", "equipment_group_id": "EG-B", "band_group": "TEST", "assigned_resource": "101.000000 MHz"},
        {"task_unit_id": "TU-C", "equipment_group_id": "EG-C", "band_group": "TEST", "assigned_resource": "99.000000 MHz"},
        {"task_unit_id": "TU-B", "equipment_group_id": "EG-D", "band_group": "TEST", "assigned_resource": "100.000000 MHz"},
    ]

    risks = _reuse_risks(assignments, task_units, equipment_groups)
    risk_types = {item["risk_type"] for item in risks}

    assert "同频冲突" in risk_types
    assert "任务频谱竞争" in risk_types
    assert "高功率近距耦合" in risk_types
    assert "三阶互调风险" in risk_types
    assert all(0 <= item["score"] <= 100 for item in risks)


def test_task_conflict_model_reports_adjacent_and_guard_shortage() -> None:
    task_units = [
        {"task_unit_id": "TU-A", "area_center_lat": 30.0, "area_center_lon": 110.0, "priority": 2},
        {"task_unit_id": "TU-B", "area_center_lat": 31.0, "area_center_lon": 111.0, "priority": 2},
    ]
    equipment_groups = [
        {"equipment_group_id": "EG-A", "task_unit_id": "TU-A", "tx_power_w": 10, "protection_distance_km": 5, "min_spacing_khz": 100, "guard_band_khz": 100, "bandwidth_khz": 25},
        {"equipment_group_id": "EG-B", "task_unit_id": "TU-B", "tx_power_w": 10, "protection_distance_km": 5, "min_spacing_khz": 100, "guard_band_khz": 100, "bandwidth_khz": 25},
    ]
    assignments = [
        {"task_unit_id": "TU-A", "equipment_group_id": "EG-A", "band_group": "TEST", "assigned_resource": "410.000000 MHz"},
        {"task_unit_id": "TU-B", "equipment_group_id": "EG-B", "band_group": "TEST", "assigned_resource": "410.050000 MHz"},
    ]

    risks = _reuse_risks(assignments, task_units, equipment_groups)
    risk_types = {item["risk_type"] for item in risks}

    assert {"邻频冲突", "保护间隔不足"} <= risk_types


def test_weighted_plan_decision_uses_continuous_six_dimension_weights() -> None:
    summary = {
        "task_satisfaction_avg": 92,
        "quality_scores": {
            "items": [
                {"key": "task_assurance", "score": 92},
                {"key": "interference_risk", "score": 35},
                {"key": "spectrum_efficiency", "score": 80},
                {"key": "executability", "score": 90},
                {"key": "change_cost", "score": 70},
            ]
        },
    }
    assurance_first = _weighted_plan_decision(summary, {"task": 100, "risk": 5, "spectrum": 5, "priority": 5, "switching": 5, "reuse": 5})
    risk_first = _weighted_plan_decision(summary, {"task": 5, "risk": 100, "spectrum": 5, "priority": 5, "switching": 5, "reuse": 5})

    assert assurance_first["score"] > risk_first["score"]
    assert len(assurance_first["components"]) == 6
    assert next(item for item in assurance_first["components"] if item["key"] == "task")["contribution"] > 70


def test_all_task_objectives_are_supported() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    for objective in [item["objective"] for item in TASK_OBJECTIVES]:
        result = solve_task_assignment(task_units, equipment_groups, spectrum_rules, objective)
        assert result["status"] == "success"
        assert result["summary"]["equipment_group_count"] == 20


def test_demo_scenario_variants_are_supported() -> None:
    keys = {item["key"] for item in TASK_SCENARIOS}
    assert {"baseline", "uav_priority", "radar_priority", "backhaul_limited", "large_joint_exercise", "stress_performance"} <= keys
    task_units, equipment_groups, spectrum_rules = demo_scenario_records("uav_priority")
    assert any(item["equipment_group_id"] == "EG-UAV-VIDEO" for item in equipment_groups)
    assert any(item["band_group"] == "S-SIM-2" for item in spectrum_rules)
    assert next(item for item in task_units if item["task_unit_id"] == "TU-UAV")["priority"] == 10


def test_scaled_demo_scenarios_increase_task_and_equipment_counts() -> None:
    large_units, large_groups, large_rules = demo_scenario_records("large_joint_exercise")
    stress_units, stress_groups, stress_rules = demo_scenario_records("stress_performance")
    large_types = {item["equipment_type"] for item in large_groups}
    assert len(large_units) >= 18
    assert len(large_groups) >= 70
    assert sum(item["count"] for item in large_groups) >= 1500
    assert len(stress_units) >= 36
    assert len(stress_groups) >= 150
    assert sum(item["count"] for item in stress_groups) >= 3500
    assert {"战术数据链终端", "宽带自组网节点", "导航授时终端", "频谱侦察接收机", "电子压制设备"} <= large_types
    assert any(item["band_group"] == "UHF-SIM-2" for item in large_rules)
    assert any(item["band_group"] == "PNT-SIM-2" for item in large_rules)
    assert any(item["band_group"] == "EW-SIM-2" for item in large_rules)
    assert any(item["rule_type"] == "保护" for item in stress_rules)


def test_task_planning_performance_test_returns_rows() -> None:
    result = task_planning_performance_test(["baseline", "large_joint_exercise"])
    assert result["summary"]["scenario_count"] == 2
    assert result["summary"]["run_count"] == len(TASK_OBJECTIVES) * 2
    assert result["summary"]["max_equipment_group_count"] >= 70
    assert result["analysis"]["leader_summary"]
    assert len(result["analysis"]["scale_curve"]) == 2
    assert len(result["analysis"]["objective_rankings"]) == 2
    assert result["analysis"]["diagnostics"]
    capacity = result["analysis"]["capacity_profile"]
    assert capacity["recommended"]["task_unit_count"] > 0
    assert capacity["largest"]["equipment_group_count"] >= capacity["recommended"]["equipment_group_count"]
    assert capacity["status"] in {"容量充足", "需分批规划", "需扩容规则"}
    assert all(row["elapsed_ms"] >= 1 for row in result["rows"])
    assert all(row["equipment_sample_count"] > 0 for row in result["rows"])


def test_task_sample_catalog_and_batch_performance_are_available() -> None:
    catalog = task_sample_catalog()
    assert len(catalog["presets"]) >= 3
    assert len(catalog["task_unit_profiles"]) >= 8
    assert len(catalog["equipment_types"]) >= 15
    assert len(catalog["weight_templates"]) >= 5
    assert catalog["parametric_defaults"]["unit_count"] >= 20
    assert catalog["objective_count"] == len(TASK_OBJECTIVES)
    result = task_planning_batch_performance_test()
    assert result["summary"]["scenario_count"] == len(catalog["batch_scales"])
    assert result["summary"]["run_count"] == len(catalog["batch_scales"]) * len(TASK_OBJECTIVES)
    assert result["summary"]["max_equipment_group_count"] >= 150
    assert result["analysis"]["leader_summary"]
    assert result["analysis"]["capacity_profile"]["largest"]["equipment_group_count"] >= 150
    assert result["analysis"]["capacity_profile"]["problem_focus"]


def test_parametric_demo_records_adjust_equipment_mix_and_rule_density() -> None:
    task_units, equipment_groups, spectrum_rules = parametric_demo_records(
        {
            "unit_count": 16,
            "density_multiplier": 1.6,
            "radar_ratio": 35,
            "uav_ratio": 25,
            "protection_density": 4,
            "forbidden_density": 3,
        }
    )
    assert len(task_units) == 16
    assert len(equipment_groups) > 50
    assert sum(item["count"] for item in equipment_groups) > 800
    assert sum(1 for item in task_units if "雷达" in item["unit_type"]) >= 4
    assert sum(1 for item in task_units if "无人机" in item["unit_type"]) >= 3
    assert sum(1 for item in spectrum_rules if item.get("source") == "SIM_PARAMETRIC" and item["rule_type"] == "保护") == 4
    assert sum(1 for item in spectrum_rules if item.get("source") == "SIM_PARAMETRIC" and item["rule_type"] == "禁用") == 3


def test_task_visualization_contains_bands_and_reasons() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    result = solve_task_assignment(task_units, equipment_groups, spectrum_rules, "task_assurance")
    visualization = build_task_visualization_data(
        task_units=task_units,
        equipment_groups=equipment_groups,
        spectrum_rules=spectrum_rules,
        assignments=result["assignments"],
        risk_items=result["risk_items"],
        summary=result["summary"],
    )
    assert len(visualization["task_units"]) == 12
    assert len(visualization["band_usage"]) >= 5
    assert "partial_reason_rank" in visualization
    assert "quality_scores" in visualization["summary"]
    assert "bottleneck_analysis" in visualization["summary"]
    assert "spectrum_contention" in visualization["summary"]
    assert visualization["summary"]["bottleneck_analysis"]["band_bottlenecks"]
    assert any("alternative_resources" in item for item in visualization["assignments"])
    assert any("explanation_chain" in item for item in visualization["assignments"])


def test_weighted_objective_forced_band_and_required_full_constraints() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    target_groups = [item for item in equipment_groups if item["equipment_group_id"] in {"EG-UAV-DATA", "EG-RAD-TRACK"}]
    result = solve_task_assignment(
        task_units,
        target_groups,
        spectrum_rules,
        "task_assurance",
        constraints={
            "weights": {"risk": 95, "task": 60, "spectrum": 40, "priority": 60, "switching": 20, "reuse": 20},
            "strategy_profile": "risk_first",
            "forced_band_groups": {"EG-UAV-DATA": "S-SIM-1"},
            "required_full_targets": ["EG-RAD-TRACK"],
            "allow_low_priority_degrade": False,
        },
    )
    assert result["summary"]["requested_objective"] == "task_assurance"
    assert result["summary"]["effective_objective"] == "minimize_interference"
    assert result["summary"]["strategy_profile"] == "risk_first"
    assert result["summary"]["constraint_weights"]["risk"] == 95
    uav = next(item for item in result["assignments"] if item["equipment_group_id"] == "EG-UAV-DATA")
    radar = next(item for item in result["assignments"] if item["equipment_group_id"] == "EG-RAD-TRACK")
    assert uav["band_group"] == "S-SIM-1"
    assert radar["required_full"] is True
    assert radar["status"] in {"完全满足", "部分满足", "未满足"}


def test_strategy_trials_return_ranked_adoptable_candidates() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    baseline = solve_task_assignment(task_units, equipment_groups, spectrum_rules, "task_assurance")
    changes = {
        "objective": "task_assurance",
        "strategy_profile": "balanced",
        "constraint_weights": {},
        "available_ranges": [],
        "forbidden_ranges": [],
        "priority_updates": [],
        "satisfaction_updates": [],
        "locked_equipment_group_ids": [],
        "locked_task_unit_ids": [],
        "avoid_band_groups": [],
        "forced_band_groups": {},
        "required_full_targets": [],
        "allow_low_priority_degrade": True,
    }

    trials = build_task_strategy_trials(
        task_units,
        equipment_groups,
        spectrum_rules,
        changes,
        base_summary=baseline["summary"],
        base_run_id=42,
    )

    assert trials["ok"] is True
    assert trials["base_run_id"] == 42
    assert len(trials["candidates"]) >= 4
    assert trials["recommended_trial_id"] == trials["candidates"][0]["trial_id"]
    assert trials["candidates"][0]["recommended"] is True
    trial_ids = {item["trial_id"] for item in trials["candidates"]}
    assert any(item.startswith("current_request__") for item in trial_ids)
    assert any(item.startswith("risk_first__") for item in trial_ids)
    assert any(item.startswith("spectrum_saving__") for item in trial_ids)
    assert any(item.startswith("task_guard__") for item in trial_ids)
    decision_summary = trials["decision_summary"]
    assert decision_summary["recommended_trial_id"] == trials["recommended_trial_id"]
    assert {item["role"] for item in decision_summary["role_picks"]} >= {
        "recommended",
        "risk_minimum",
        "satisfaction_maximum",
        "bandwidth_minimum",
        "quality_maximum",
    }
    assert decision_summary["pareto_frontier"]
    assert decision_summary["tradeoff_notes"]
    assert decision_summary["adoption_guardrails"]
    for candidate in trials["candidates"]:
        assert "quality_total" in candidate
        assert "deltas" in candidate
        assert "constraint_variant" in candidate
        assert candidate["constraint_label"]
        assert candidate["constraint_changes"]
        assert candidate["apply_payload"]["base_run_id"] == 42
        assert candidate["apply_payload"]["objective"] == candidate["objective"]
        assert candidate["apply_payload"]["strategy_profile"] == candidate["strategy_profile"]


def test_strategy_trials_include_constraint_variants_for_shortfalls() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    baseline = solve_task_assignment(task_units, equipment_groups, spectrum_rules, "task_assurance")
    visualization = build_task_visualization_data(
        task_units=task_units,
        equipment_groups=equipment_groups,
        spectrum_rules=spectrum_rules,
        assignments=baseline["assignments"],
        risk_items=baseline["risk_items"],
        summary=baseline["summary"],
    )
    changes = {
        "objective": "task_assurance",
        "strategy_profile": "balanced",
        "constraint_weights": {},
        "available_ranges": [],
        "forbidden_ranges": [],
        "priority_updates": [],
        "satisfaction_updates": [],
        "locked_equipment_group_ids": [],
        "locked_task_unit_ids": [],
        "avoid_band_groups": [],
        "forced_band_groups": {},
        "required_full_targets": [],
        "allow_low_priority_degrade": True,
    }

    variants = _strategy_constraint_variants(changes, baseline["summary"], baseline["assignments"], visualization)
    variant_ids = {item["variant_id"] for item in variants}
    assert "current_constraints" in variant_ids
    assert "required_full_shortfalls" in variant_ids
    assert any(item.get("overrides", {}).get("available_ranges") for item in variants)

    trials = build_task_strategy_trials(
        task_units,
        equipment_groups,
        spectrum_rules,
        changes,
        base_summary=baseline["summary"],
        base_run_id=42,
        constraint_variants=variants,
    )

    assert trials["ok"] is True
    assert any(candidate["constraint_variant"] != "current_constraints" for candidate in trials["candidates"])
    assert any(
        candidate["apply_payload"].get("required_full_targets") or candidate["apply_payload"].get("available_ranges")
        for candidate in trials["candidates"]
    )


def test_trial_context_is_sanitized_for_audit_and_preview_items() -> None:
    context = _sanitize_trial_context(
        {
            "trial_id": "task_guard__supplement",
            "label": "任务保障 + 补频",
            "objective_label": "任务保障优先",
            "constraint_label": "补充短板频段",
            "decision": "建议采用",
            "recommendation_score": "88.2",
            "expected_deltas": {"quality_total": "4.5", "high_risk_count": -1, "ignored": 99},
            "expected_metrics": {"task_satisfaction_avg": 94.2, "bad": "x"},
            "constraint_changes": ["补充 S-SIM-2", "", "要求 TU-RADAR 完全满足"],
            "raw_rows": [{"should": "drop"}],
        }
    )
    assert context["trial_id"] == "task_guard__supplement"
    assert context["recommendation_score"] == 88.2
    assert context["expected_deltas"] == {"quality_total": 4.5, "high_risk_count": -1.0}
    assert context["expected_metrics"] == {"task_satisfaction_avg": 94.2}
    assert "raw_rows" not in context

    items = _replan_change_items(
        {
            "trial_context": context,
            "reuse_strategy_from_run_id": 0,
            "objective_changed": False,
            "available_ranges": [],
            "forbidden_ranges": [],
            "priority_updates": [],
            "satisfaction_updates": [],
            "locked_equipment_group_ids": [],
            "locked_task_unit_ids": [],
            "avoid_band_groups": [],
            "forced_band_groups": {},
            "required_full_targets": [],
            "constraint_weights_changed": False,
        }
    )
    assert items[0]["type"] == "采用试算候选"
    assert items[0]["target"] == "任务保障 + 补频"


def test_trial_prediction_alignment_compares_expected_and_actual_directions() -> None:
    alignment = _trial_prediction_alignment(
        {"quality_total": 4, "high_risk_count": -1, "task_satisfaction_avg": 0},
        {"quality_total": 3, "high_risk_count": 1, "task_satisfaction_avg": 0},
    )
    assert alignment["total_count"] == 3
    assert alignment["matched_count"] == 2
    assert "2/3" in alignment["summary"]
    rows = {item["key"]: item for item in alignment["rows"]}
    assert rows["quality_total"]["direction_matched"] is True
    assert rows["high_risk_count"]["direction_matched"] is False


def test_task_agent_assessment_builds_maturity_gates_and_actions() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    result = solve_task_assignment(task_units, equipment_groups, spectrum_rules, "task_assurance")
    visualization = build_task_visualization_data(
        task_units=task_units,
        equipment_groups=equipment_groups,
        spectrum_rules=spectrum_rules,
        assignments=result["assignments"],
        risk_items=result["risk_items"],
        summary=result["summary"],
    )
    assessment = build_task_agent_assessment(
        result["summary"],
        result["assignments"],
        result["risk_items"],
        visualization,
        versions={
            "runs": [{"run_id": 1}, {"run_id": 2}],
            "audit_logs": [{"action": "task_replan"}, {"action": "task_plan"}],
            "performance_history": [{"run_count": 20}],
        },
        run={"run_id": 2, "elapsed_ms": 120},
    )
    assert assessment["status"] == "ready"
    assert assessment["maturity_score"] > 0
    assert len(assessment["capability_items"]) >= 8
    assert any(item["name"] == "任务保障率达到可执行阈值" for item in assessment["acceptance_gates"])
    assert assessment["next_actions"]
    assert len(assessment["strategy_history"]) == 2
    assert any(item["item"] == "解释链" and item["covered"] for item in assessment["artifact_checklist"])


def test_replan_preview_is_not_counted_as_executed_replan() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    result = solve_task_assignment(task_units, equipment_groups, spectrum_rules, "task_assurance")
    visualization = build_task_visualization_data(
        task_units=task_units,
        equipment_groups=equipment_groups,
        spectrum_rules=spectrum_rules,
        assignments=result["assignments"],
        risk_items=result["risk_items"],
        summary=result["summary"],
    )
    assessment = build_task_agent_assessment(
        result["summary"],
        result["assignments"],
        result["risk_items"],
        visualization,
        versions={
            "runs": [{"run_id": 1}],
            "audit_logs": [{"action": "task_replan_preview"}],
            "performance_history": [],
        },
        run={"run_id": 1, "elapsed_ms": 120},
    )
    dynamic_item = next(item for item in assessment["capability_items"] if item["key"] == "dynamic_replanning")
    assert "1 次重规划预览" in dynamic_item["evidence"]
    assert "0 次执行重规划" in dynamic_item["evidence"]


def test_agent_assessment_suggests_capacity_boundary_split() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    result = solve_task_assignment(task_units, equipment_groups, spectrum_rules, "task_assurance")
    visualization = build_task_visualization_data(
        task_units=task_units,
        equipment_groups=equipment_groups,
        spectrum_rules=spectrum_rules,
        assignments=result["assignments"],
        risk_items=result["risk_items"],
        summary=result["summary"],
    )
    capacity_profile = {
        "status": "需分批规划",
        "problem_focus": "可用窗口或连续带宽不足",
        "recommended": {"task_unit_count": 12, "equipment_group_count": 52},
        "largest": {"task_unit_count": 36, "equipment_group_count": 153},
    }
    assessment = build_task_agent_assessment(
        result["summary"],
        result["assignments"],
        result["risk_items"],
        visualization,
        versions={
            "runs": [{"run_id": 1}, {"run_id": 2}],
            "audit_logs": [{"action": "task_replan"}],
            "performance_history": [{"run_count": 50, "capacity_profile": capacity_profile}],
        },
        run={"run_id": 1, "elapsed_ms": 120},
    )
    action = next(item for item in assessment["next_actions"] if item["action_id"] == "capacity_boundary_split")
    assert action["deterministic_payload"]["operation"] == "capacity_batch_planning"
    assert action["deterministic_payload"]["capacity_profile"]["recommended"]["task_unit_count"] == 12
    assert assessment["scale_evidence"]["status"] == "需分批规划"
    assert assessment["scale_evidence"]["capacity_defined"] is True
    assert "推荐单次 12 单元/52 装备组" in assessment["scale_evidence"]["evidence"]
    gates = {item["name"]: item for item in assessment["acceptance_gates"]}
    assert gates["具备规模压测证据"]["passed"] is True
    assert gates["容量边界可执行"]["passed"] is True
    artifact = next(item for item in assessment["artifact_checklist"] if item["item"] == "规模效能证据")
    assert artifact["covered"] is True
    scalability = next(item for item in assessment["capability_items"] if item["key"] == "scalability")
    assert "12" in scalability["evidence"]


def test_agent_assessment_builds_capacity_batch_plan_when_current_exceeds_boundary() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    result = solve_task_assignment(task_units, equipment_groups, spectrum_rules, "task_assurance")
    visualization = build_task_visualization_data(
        task_units=task_units,
        equipment_groups=equipment_groups,
        spectrum_rules=spectrum_rules,
        assignments=result["assignments"],
        risk_items=result["risk_items"],
        summary=result["summary"],
    )
    capacity_profile = {
        "status": "需分批规划",
        "problem_focus": "当前任务单元数量超过推荐单次边界",
        "recommended": {"task_unit_count": 2, "equipment_group_count": 99, "equipment_sample_count": 1000},
        "largest": {"task_unit_count": 6, "equipment_group_count": 120, "equipment_sample_count": 1800},
    }
    assessment = build_task_agent_assessment(
        result["summary"],
        result["assignments"],
        result["risk_items"],
        visualization,
        versions={
            "runs": [{"run_id": 1}, {"run_id": 2}],
            "audit_logs": [{"action": "task_replan"}],
            "performance_history": [{"run_count": 50, "capacity_profile": capacity_profile}],
        },
        run={"run_id": 2, "elapsed_ms": 120},
    )

    batch_plan = assessment["scale_evidence"]["batch_plan"]
    assert batch_plan["available"] is True
    assert batch_plan["needed"] is True
    assert batch_plan["batch_count"] >= 2
    assert all(batch["task_unit_count"] <= 2 for batch in batch_plan["batches"])
    action = next(item for item in assessment["next_actions"] if item["action_id"] == "capacity_boundary_split")
    assert action["deterministic_payload"]["batch_plan"]["batch_count"] == batch_plan["batch_count"]
    assert "当前预案" in action["suggested_message"]


def test_capacity_batch_execution_creates_scoped_sub_runs() -> None:
    engine = create_engine("sqlite:///:memory:")
    SQLModel.metadata.create_all(engine)
    task_units, equipment_groups, spectrum_rules = demo_scenario_records("large_joint_exercise")
    with Session(engine) as session:
        project = Project(name="capacity-batch-test")
        session.add(project)
        session.commit()
        session.refresh(project)
        assert project.id is not None

        replace_task_units(session, project.id, task_units)
        replace_equipment_groups(session, project.id, equipment_groups)
        replace_spectrum_rules(session, project.id, spectrum_rules)
        session.add(
            AuditLog(
                project_id=project.id,
                actor="system",
                action="task_performance_batch",
                detail=json.dumps(
                    {
                        "summary": {
                            "scenario_count": 5,
                            "run_count": 50,
                            "max_elapsed_ms": 80,
                            "avg_elapsed_ms": 30,
                            "best_quality_total": 82,
                        },
                        "leader_summary": "capacity boundary fixture",
                        "capacity_profile": {
                            "status": "需分批规划",
                            "problem_focus": "fixture",
                            "recommended": {"task_unit_count": 4, "equipment_group_count": 99, "equipment_sample_count": 1200},
                            "largest": {"task_unit_count": 18, "equipment_group_count": 80, "equipment_sample_count": 1800},
                        },
                    },
                    ensure_ascii=False,
                ),
            )
        )
        session.commit()

        base_run = run_task_planning(session, project.id, "task_assurance")
        result = execute_capacity_batch_planning(session, project.id, {"base_run_id": base_run.id})

        assert result["ok"] is True
        assert result["base_run_id"] == base_run.id
        assert len(result["runs"]) >= 2
        assert result["cross_batch_review"]["summary"]
        merged = result["merged_plan"]
        assert merged["base_run_id"] == base_run.id
        assert merged["batch_count"] == len(result["runs"])
        assert merged["task_unit_count"] == sum(row["task_unit_count"] for row in result["runs"])
        assert merged["equipment_group_count"] == sum(row["equipment_group_count"] for row in result["runs"])
        assert set(merged["run_ids"]) == {row["run_id"] for row in result["runs"]}
        assert merged["unresolved_count"] >= merged["cross_batch_risk_count"]
        assert merged["audit_items"]
        assert merged["recommendation"]
        run_ids = [item["run_id"] for item in result["runs"]]
        assert len(run_ids) == len(set(run_ids))
        for row in result["runs"]:
            summary = row["summary"]
            assert summary["batch_context"]["base_run_id"] == base_run.id
            assert summary["task_unit_count"] <= 4
            assignments = session.exec(select(EquipmentAssignment).where(EquipmentAssignment.run_id == row["run_id"])).all()
            assert len(assignments) == summary["equipment_group_count"]
            assert {item.task_unit_id for item in assignments} <= set(summary["batch_context"]["task_unit_ids"])

        export_bytes, audit_id = latest_capacity_batch_export_xlsx(session, project.id)
        assert audit_id > 0
        workbook = load_workbook(BytesIO(export_bytes), read_only=True)
        assert {
            "合并总方案",
            "子规划摘要",
            "跨批风险",
            "审核清单",
            "合并指配明细",
            "批内风险明细",
        } <= set(workbook.sheetnames)
        assert workbook["合并总方案"].max_row >= 2
        assert workbook["子规划摘要"].max_row == len(result["runs"]) + 1

        closure = execute_capacity_batch_risk_closure(session, project.id, {})
        assert closure["ok"] is True
        assert closure["base_run_id"] == base_run.id
        assert closure["closure"]["applied_rule_count"] > 0
        assert closure["closure"]["before_cross_batch_risk_count"] > 0
        assert "after_cross_batch_risk_count" in closure["closure"]
        assert closure["merged_plan"]["run_ids"]
        assert len(closure["runs"]) == len(result["runs"])
        for row in closure["runs"]:
            summary = row["summary"]
            assert summary["batch_context"]["base_run_id"] == base_run.id
            assert summary["risk_closure_context"]["applied_rule_count"] == closure["closure"]["applied_rule_count"]
            assert "hard_avoid_rule_count" in summary

        closure_export_bytes, closure_audit_id = latest_capacity_batch_closure_export_xlsx(session, project.id)
        assert closure_audit_id > audit_id
        closure_workbook = load_workbook(BytesIO(closure_export_bytes), read_only=True)
        assert {
            "闭环摘要",
            "闭环前后对比",
            "闭环硬约束",
            "闭环合并总方案",
            "闭环子规划摘要",
            "闭环后跨批风险",
            "闭环指配明细",
            "闭环批内风险",
        } <= set(closure_workbook.sheetnames)
        assert closure_workbook["闭环前后对比"].max_row >= 5
        assert closure_workbook["闭环硬约束"].max_row == closure["closure"]["applied_rule_count"] + 1

        versions = task_versions_and_audit(session, project.id)
        assert versions["closure_events"]
        event = versions["closure_events"][0]
        assert event["id"] == closure_audit_id
        assert event["applied_rule_count"] == closure["closure"]["applied_rule_count"]
        assert event["before_cross_batch_risk_count"] == closure["closure"]["before_cross_batch_risk_count"]
        assert event["after_cross_batch_risk_count"] == closure["closure"]["after_cross_batch_risk_count"]
        assert event["run_ids"] == [row["run_id"] for row in closure["runs"]]
        assert event["rules"]


def test_replan_change_items_omit_unchanged_objective_and_weights() -> None:
    items = _replan_change_items(
        {
            "objective": "task_assurance",
            "objective_changed": False,
            "available_ranges": [{"band_group": "C-SIM-1", "start_mhz": 4501, "end_mhz": 4531}],
            "forbidden_ranges": [],
            "priority_updates": [],
            "satisfaction_updates": [],
            "locked_equipment_group_ids": [],
            "locked_task_unit_ids": [],
            "avoid_band_groups": [],
            "forced_band_groups": {},
            "required_full_targets": [],
            "constraint_weights": {"task": 70, "risk": 70},
            "constraint_weights_changed": False,
            "strategy_profile": "balanced",
        }
    )
    assert [item["type"] for item in items] == ["新增可用频段"]


def test_replan_change_items_include_reused_strategy() -> None:
    items = _replan_change_items(
        {
            "objective": "minimize_interference",
            "reuse_strategy_from_run_id": 7,
            "objective_changed": True,
            "available_ranges": [],
            "forbidden_ranges": [],
            "priority_updates": [],
            "satisfaction_updates": [],
            "locked_equipment_group_ids": [],
            "locked_task_unit_ids": [],
            "avoid_band_groups": [],
            "forced_band_groups": {},
            "required_full_targets": [],
            "constraint_weights": {"risk": 95},
            "constraint_weights_changed": True,
            "strategy_profile": "risk_first",
        }
    )
    assert items[0]["type"] == "复用历史策略"
    assert items[0]["target"] == "#7"
    assert "干扰最低" in items[0]["detail"]


def test_solver_respects_band_compatibility_and_power_limits() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    radar_group = [item for item in equipment_groups if item["equipment_group_id"] == "EG-RAD-LOW"]
    result = solve_task_assignment(task_units, radar_group, spectrum_rules, "task_assurance")
    assignment = result["assignments"][0]
    assert assignment["band_group"] == "C-SIM-1"
    assert assignment["status"] in {"完全满足", "部分满足"}


def test_duplex_pair_keeps_minimum_spacing() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    relay_group = [item for item in equipment_groups if item["equipment_group_id"] == "EG-MOB-RELAY"]
    result = solve_task_assignment(task_units, relay_group, spectrum_rules, "task_assurance")
    assignment = result["assignments"][0]
    segments = _parse_resource_segments(assignment["assigned_resource"])
    assert len(segments) >= 2
    assert segments[1].start - segments[0].start >= 0.099


def test_replan_message_extracts_forbidden_range() -> None:
    _, _, spectrum_rules = demo_scenario_records()
    ranges = _forbidden_ranges_from_message("禁用 2210-2215 MHz，尽量少用 C-SIM-1", spectrum_rules)
    assert ranges == [
        {
            "band_group": "S-SIM-1",
            "start_mhz": 2210.0,
            "end_mhz": 2215.0,
            "reason": "用户聊天追加禁用",
        }
    ]


def test_replan_message_extracts_available_range() -> None:
    _, _, spectrum_rules = demo_scenario_records()
    ranges = _available_ranges_from_message("补充可用 S-SIM-1 2235-2240 MHz", spectrum_rules)
    assert ranges == [
        {
            "band_group": "S-SIM-1",
            "start_mhz": 2235.0,
            "end_mhz": 2240.0,
            "reason": "用户聊天补充可用频段",
        }
    ]


def test_agent_action_can_suggest_resource_supplement_payload() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records("stress_performance")
    result = solve_task_assignment(task_units, equipment_groups, spectrum_rules, "task_assurance")
    visualization = build_task_visualization_data(
        task_units=task_units,
        equipment_groups=equipment_groups,
        spectrum_rules=spectrum_rules,
        assignments=result["assignments"],
        risk_items=result["risk_items"],
        summary=result["summary"],
    )
    assessment = build_task_agent_assessment(
        result["summary"],
        result["assignments"],
        result["risk_items"],
        visualization,
        versions={"runs": [{"run_id": 1}], "audit_logs": [], "performance_history": []},
        run={"run_id": 1, "elapsed_ms": 120},
    )
    supplement_actions = [item for item in assessment["next_actions"] if item["deterministic_payload"].get("available_ranges")]
    assert supplement_actions
    available_range = supplement_actions[0]["deterministic_payload"]["available_ranges"][0]
    assert available_range["band_group"]
    assert available_range["end_mhz"] > available_range["start_mhz"]


def test_partial_action_adds_supplement_window_and_required_targets() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    result = solve_task_assignment(task_units, equipment_groups, spectrum_rules, "task_assurance")
    visualization = build_task_visualization_data(
        task_units=task_units,
        equipment_groups=equipment_groups,
        spectrum_rules=spectrum_rules,
        assignments=result["assignments"],
        risk_items=result["risk_items"],
        summary=result["summary"],
    )
    assessment = build_task_agent_assessment(
        result["summary"],
        result["assignments"],
        result["risk_items"],
        visualization,
        versions={"runs": [{"run_id": 1}, {"run_id": 2}], "audit_logs": [{"action": "task_replan"}], "performance_history": []},
        run={"run_id": 1, "elapsed_ms": 120},
    )
    action = next(item for item in assessment["next_actions"] if item["action_id"] == "explain_or_upgrade_partial")
    payload = action["deterministic_payload"]
    assert payload["required_full_targets"]
    assert payload["available_ranges"]
    available_range = payload["available_ranges"][0]
    assert available_range["end_mhz"] > available_range["start_mhz"]
    assert available_range["end_mhz"] - available_range["start_mhz"] <= 60
    assert payload["constraint_weights"]["task"] == 95


def test_supplement_range_prefers_unmet_assignment_gap() -> None:
    assignments = [
        {
            "equipment_group_id": "EG-RAD-GAP",
            "status": "部分满足",
            "band_group": "C-SIM-1",
            "requested_channels": 4,
            "assigned_channels": 1,
        }
    ]
    visualization = {
        "spectrum_timeline": [{"band_group": "C-SIM-1", "start_mhz": 4400.0, "end_mhz": 4500.0}],
        "assignments": [
            {
                **assignments[0],
                "priority": 9,
                "bandwidth_khz": 5000,
                "explanation_chain": {"required_extra_resource": {"estimated_extra_width_mhz": 15.0}},
            }
        ],
    }
    supplement = _supplement_range_for_unmet_assignment(assignments, visualization)
    assert supplement is not None
    assert supplement["band_group"] == "C-SIM-1"
    assert supplement["start_mhz"] > 4500
    assert supplement["end_mhz"] - supplement["start_mhz"] >= 15
    assert "EG-RAD-GAP" in supplement["reason"]


def test_supplement_ranges_cover_multiple_partial_bands() -> None:
    assignments = [
        {
            "equipment_group_id": "EG-RAD-C",
            "status": "部分满足",
            "band_group": "C-SIM-1",
            "requested_channels": 4,
            "assigned_channels": 2,
        },
        {
            "equipment_group_id": "EG-RAD-S",
            "status": "部分满足",
            "band_group": "S-SIM-1",
            "requested_channels": 2,
            "assigned_channels": 1,
        },
    ]
    visualization = {
        "spectrum_timeline": [
            {"band_group": "C-SIM-1", "start_mhz": 4400.0, "end_mhz": 4500.0},
            {"band_group": "S-SIM-1", "start_mhz": 2200.0, "end_mhz": 2230.0},
        ],
        "assignments": [
            {
                **assignments[0],
                "priority": 9,
                "bandwidth_khz": 15000,
                "explanation_chain": {"required_extra_resource": {"estimated_extra_width_mhz": 30.0}},
            },
            {
                **assignments[1],
                "priority": 8,
                "bandwidth_khz": 10000,
                "explanation_chain": {"required_extra_resource": {"estimated_extra_width_mhz": 10.0}},
            },
        ],
    }
    ranges = _supplement_ranges_for_unmet_assignments(assignments, visualization)
    assert {item["band_group"] for item in ranges} == {"C-SIM-1", "S-SIM-1"}
    assert all(item["end_mhz"] > item["start_mhz"] for item in ranges)
    assert all(item["end_mhz"] - item["start_mhz"] <= 60 for item in ranges)


def test_assessment_includes_flat_replan_effect_action() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    result = solve_task_assignment(task_units, equipment_groups, spectrum_rules, "task_assurance")
    visualization = build_task_visualization_data(
        task_units=task_units,
        equipment_groups=equipment_groups,
        spectrum_rules=spectrum_rules,
        assignments=result["assignments"],
        risk_items=result["risk_items"],
        summary=result["summary"],
    )
    assessment = build_task_agent_assessment(
        result["summary"],
        result["assignments"],
        result["risk_items"],
        visualization,
        versions={"runs": [{"run_id": 1}, {"run_id": 2}], "audit_logs": [{"action": "task_replan"}], "performance_history": []},
        run={"run_id": 2, "elapsed_ms": 120},
        replan_effect={
            "status": "flat",
            "summary": "本次重规划收益不明显",
            "recommendation": "继续定位未满足对象",
        },
    )
    assert assessment["replan_effect"]["status"] == "flat"
    assert any(item["action_id"] == "explain_flat_replan" for item in assessment["next_actions"])
    assert all("rank_score" in item and "expected_gain" in item and "guardrail" in item for item in assessment["next_actions"])


def test_flat_unused_supplement_pivots_without_repeating_available_ranges() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    result = solve_task_assignment(task_units, equipment_groups, spectrum_rules, "task_assurance")
    visualization = build_task_visualization_data(
        task_units=task_units,
        equipment_groups=equipment_groups,
        spectrum_rules=spectrum_rules,
        assignments=result["assignments"],
        risk_items=result["risk_items"],
        summary=result["summary"],
    )
    assessment = build_task_agent_assessment(
        result["summary"],
        result["assignments"],
        result["risk_items"],
        visualization,
        versions={"runs": [{"run_id": 1}, {"run_id": 2}], "audit_logs": [{"action": "task_replan"}], "performance_history": []},
        run={"run_id": 2, "elapsed_ms": 120},
        replan_effect={
            "status": "flat",
            "summary": "本次重规划收益不明显",
            "recommendation": "转向风险或频谱策略",
            "added_ranges": [{"band_group": "C-SIM-1", "start_mhz": 4500, "end_mhz": 4560}],
            "added_range_usage": [{"band_group": "C-SIM-1", "used_width_mhz": 0, "used_assignment_count": 0}],
        },
    )
    pivot = next(item for item in assessment["next_actions"] if item["action_id"] == "pivot_after_flat_replan")
    assert pivot["rank_score"] >= 100
    assert not pivot["deterministic_payload"].get("available_ranges")
    assert pivot["deterministic_payload"]["objective"] in {"minimize_interference", "minimize_bandwidth"}
    repeated_supplements = [
        item
        for item in assessment["next_actions"]
        if item["action_id"] in {"explain_or_upgrade_partial", "relieve_bottleneck_band"}
        and item["deterministic_payload"].get("available_ranges")
    ]
    assert repeated_supplements == []


def test_flat_risk_first_strategy_pivots_to_spectrum_saving() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    result = solve_task_assignment(task_units, equipment_groups, spectrum_rules, "minimize_interference")
    visualization = build_task_visualization_data(
        task_units=task_units,
        equipment_groups=equipment_groups,
        spectrum_rules=spectrum_rules,
        assignments=result["assignments"],
        risk_items=result["risk_items"],
        summary=result["summary"],
    )
    assessment = build_task_agent_assessment(
        result["summary"],
        result["assignments"],
        result["risk_items"],
        visualization,
        versions={"runs": [{"run_id": 1}, {"run_id": 2}], "audit_logs": [{"action": "task_replan"}], "performance_history": []},
        run={"run_id": 2, "elapsed_ms": 120},
        replan_effect={
            "status": "flat",
            "summary": "风险优先重规划收益不明显",
            "recommendation": "转向频谱节约",
            "current_objective": "minimize_interference",
            "current_effective_objective": "minimize_interference",
            "current_strategy_profile": "risk_first",
        },
    )
    pivot = next(item for item in assessment["next_actions"] if item["action_id"] == "pivot_after_flat_replan")
    assert pivot["deterministic_payload"]["objective"] == "minimize_bandwidth"
    assert pivot["deterministic_payload"]["strategy_profile"] == "spectrum_saving"
    assert pivot["rank_score"] >= 100


def test_regressed_spectrum_saving_suggests_rollback_and_capacity_review() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    result = solve_task_assignment(task_units, equipment_groups, spectrum_rules, "minimize_bandwidth")
    visualization = build_task_visualization_data(
        task_units=task_units,
        equipment_groups=equipment_groups,
        spectrum_rules=spectrum_rules,
        assignments=result["assignments"],
        risk_items=result["risk_items"],
        summary=result["summary"],
    )
    assessment = build_task_agent_assessment(
        result["summary"],
        result["assignments"],
        result["risk_items"],
        visualization,
        versions={"runs": [{"run_id": 9}, {"run_id": 8}], "audit_logs": [{"action": "task_replan"}], "performance_history": []},
        run={"run_id": 9, "elapsed_ms": 120},
        replan_effect={
            "status": "regressed",
            "base_run_id": 8,
            "summary": "频谱节约重规划效果下降",
            "recommendation": "回退并复测容量",
            "current_objective": "minimize_bandwidth",
            "current_effective_objective": "minimize_bandwidth",
            "current_strategy_profile": "spectrum_saving",
        },
    )
    assert assessment["next_actions"][0]["action_id"] == "rollback_to_base_run"
    capacity_action = next(item for item in assessment["next_actions"] if item["action_id"] == "capacity_after_spectrum_regression")
    assert capacity_action["deterministic_payload"]["operation"] == "task_performance_batch"


def test_clean_plan_does_not_suggest_bottleneck_supplement() -> None:
    summary = {
        "unsatisfied_group_count": 0,
        "partial_group_count": 0,
        "high_risk_count": 0,
        "medium_risk_count": 0,
        "bottleneck_analysis": {
            "band_bottlenecks": [
                {"band_group": "C-SIM-1", "pressure_score": 25, "clean_available_width_mhz": 40}
            ]
        },
    }
    actions = _agent_next_actions(
        summary=summary,
        assignments=[{"equipment_group_id": "EG-CLEAN", "status": "完全满足", "band_group": "C-SIM-1"}],
        risk_items=[],
        visualization={"spectrum_timeline": [{"band_group": "C-SIM-1", "start_mhz": 4400, "end_mhz": 4500}]},
        versions={"runs": [{"run_id": 1}, {"run_id": 2}], "audit_logs": [{"action": "task_replan"}], "performance_history": [{"run_count": 10}]},
        scores={"dynamic_replanning": 90},
    )
    assert all(item["action_id"] != "relieve_bottleneck_band" for item in actions)


def test_strategy_history_marks_transition_status() -> None:
    history = _agent_strategy_history(
        {
            "runs": [
                {
                    "run_id": 3,
                    "objective": "minimize_bandwidth",
                    "elapsed_ms": 20,
                    "summary": {
                        "requested_objective": "minimize_bandwidth",
                        "effective_objective": "minimize_bandwidth",
                        "strategy_profile": "spectrum_saving",
                        "task_satisfaction_avg": 100,
                        "quality_scores": {"total": 91},
                        "high_risk_count": 0,
                        "unsatisfied_group_count": 0,
                    },
                },
                {
                    "run_id": 2,
                    "objective": "minimize_interference",
                    "elapsed_ms": 20,
                    "summary": {
                        "requested_objective": "minimize_interference",
                        "effective_objective": "minimize_interference",
                        "strategy_profile": "risk_first",
                        "task_satisfaction_avg": 90,
                        "quality_scores": {"total": 80},
                        "high_risk_count": 1,
                        "unsatisfied_group_count": 1,
                    },
                },
            ]
        }
    )
    assert [item["run_id"] for item in history] == [2, 3]
    assert history[0]["transition_status"] == "baseline"
    assert history[1]["transition_status"] == "improved"
    assert history[1]["delta_quality"] == 11


def test_regressed_replan_effect_prioritizes_rollback_action() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    result = solve_task_assignment(task_units, equipment_groups, spectrum_rules, "task_assurance")
    visualization = build_task_visualization_data(
        task_units=task_units,
        equipment_groups=equipment_groups,
        spectrum_rules=spectrum_rules,
        assignments=result["assignments"],
        risk_items=result["risk_items"],
        summary=result["summary"],
    )
    assessment = build_task_agent_assessment(
        result["summary"],
        result["assignments"],
        result["risk_items"],
        visualization,
        versions={"runs": [{"run_id": 9}, {"run_id": 8}], "audit_logs": [{"action": "task_replan"}], "performance_history": []},
        run={"run_id": 9, "elapsed_ms": 120},
        replan_effect={
            "status": "regressed",
            "base_run_id": 8,
            "summary": "本次重规划效果下降",
            "recommendation": "回退到上一版",
        },
    )
    first = assessment["next_actions"][0]
    assert first["action_id"] == "rollback_to_base_run"
    assert first["deterministic_payload"] == {"operation": "select_task_run", "run_id": 8}
    assert first["rank_score"] > 100


def test_locked_assignment_keeps_current_resource() -> None:
    task_units, equipment_groups, spectrum_rules = demo_scenario_records()
    first = solve_task_assignment(task_units, equipment_groups, spectrum_rules, "task_assurance")
    locked = {item["equipment_group_id"]: item for item in first["assignments"] if item["equipment_group_id"] == "EG-CMD-BASE"}
    second = solve_task_assignment(
        task_units,
        equipment_groups,
        spectrum_rules,
        "priority_equipment",
        constraints={"locked_assignments": locked, "avoid_band_groups": ["UHF-SIM-1"]},
    )
    before = locked["EG-CMD-BASE"]["assigned_resource"]
    after = next(item for item in second["assignments"] if item["equipment_group_id"] == "EG-CMD-BASE")["assigned_resource"]
    assert after == before
