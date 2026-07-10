from __future__ import annotations

import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.pool import StaticPool
from sqlmodel import Session, SQLModel, create_engine

from backend.app.database import get_session
from backend.app.main import app
from backend.app.services.equipment_planning import demo_scenario_records, solve_task_assignment, validate_task_inputs
from backend.app.services.excel_io import (
    parse_equipment_group_excel,
    parse_spectrum_rule_excel,
    parse_task_unit_excel,
)
from backend.app.services.sim_scenarios import SIM_SCENARIO_KEYS, SIM_SCENARIO_MANIFESTS, sim_scenario_records
from backend.app.services.task_catalog import TASK_SCENARIOS
from scripts.build_sim_scenario_templates import build_templates


ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def api_client() -> TestClient:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    SQLModel.metadata.create_all(engine)

    def override_get_session():
        with Session(engine) as session:
            yield session

    app.dependency_overrides[get_session] = override_get_session
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.clear()
        engine.dispose()


def _create_project(client: TestClient, name: str) -> int:
    response = client.post("/api/projects", json={"name": name})
    assert response.status_code == 200, response.text
    return int(response.json()["id"])


def _upload(client: TestClient, project_id: int, endpoint: str, path: Path) -> dict:
    with path.open("rb") as handle:
        response = client.post(
            f"/api/projects/{project_id}/{endpoint}",
            files={"file": (path.name, handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert response.status_code == 200, response.text
    return response.json()


def test_sim_scenarios_are_exposed_in_the_catalog() -> None:
    catalog_keys = {item["key"] for item in TASK_SCENARIOS}
    assert SIM_SCENARIO_KEYS <= catalog_keys
    assert len(SIM_SCENARIO_KEYS) == 4


@pytest.mark.parametrize("scenario", sorted(SIM_SCENARIO_KEYS))
def test_sim_scenario_records_are_valid_synthetic_inputs(scenario: str) -> None:
    task_units, equipment_groups, spectrum_rules = sim_scenario_records(scenario)

    assert demo_scenario_records(scenario) == (task_units, equipment_groups, spectrum_rules)
    validation = validate_task_inputs(task_units, equipment_groups, spectrum_rules)
    assert validation["ok"] is True, validation
    assert validation["warnings"] == []
    assert all(item["task_unit_id"].startswith("SIM-") for item in task_units)
    assert all(item["equipment_group_id"].startswith("SIM-") for item in equipment_groups)
    assert all(item["rule_id"].startswith("SIM-") for item in spectrum_rules)
    assert all("-SIM-" in item["band_group"] for item in spectrum_rules)
    assert all(item["source"] == "SIM_CASE_TEMPLATE" for item in spectrum_rules)

    operational_payload = json.dumps(
        {"task_units": task_units, "equipment_groups": equipment_groups, "spectrum_rules": spectrum_rules},
        ensure_ascii=False,
    )
    for forbidden in ("http://", "https://", "CASE-", "PAR-", "SRC-"):
        assert forbidden not in operational_payload

    plan = solve_task_assignment(task_units, equipment_groups, spectrum_rules, "task_assurance")
    assert plan["status"] == "success"
    assert len(plan["assignments"]) == len(equipment_groups)
    assert all(item["status"] == "完全满足" for item in plan["assignments"])
    assert all(item["satisfaction_ratio"] == 1.0 for item in plan["assignments"])


def test_sim_manifests_keep_public_references_separate_from_operational_inputs() -> None:
    for scenario, manifest in SIM_SCENARIO_MANIFESTS.items():
        assert scenario in SIM_SCENARIO_KEYS
        assert manifest["scenario_id"].startswith("SIM-")
        assert manifest["reference_case_id"].startswith("CASE-")
        assert manifest["source_ids"]
        assert manifest["simulation_assumptions"]
        assert manifest["dynamic_events"]
        assert manifest["recommended_objectives"]


def test_generated_workbooks_round_trip_through_platform_parsers(tmp_path: Path) -> None:
    generated = build_templates(tmp_path)
    assert len(generated) == 16

    parsers = {
        "task_units.xlsx": parse_task_unit_excel,
        "equipment_groups.xlsx": parse_equipment_group_excel,
        "spectrum_rules.xlsx": parse_spectrum_rule_excel,
    }
    for scenario in sorted(SIM_SCENARIO_KEYS):
        scenario_dir = tmp_path / scenario
        manifest = json.loads((scenario_dir / "manifest.json").read_text(encoding="utf-8"))
        expected_counts = manifest["counts"]
        for filename, parser in parsers.items():
            path = scenario_dir / filename
            parsed = parser(path.read_bytes())
            count_key = filename.removesuffix(".xlsx")
            assert parsed.missing_columns == []
            assert len(parsed.records) == expected_counts[count_key]

            workbook = load_workbook(path, data_only=False)
            sheet = workbook.active
            assert sheet.freeze_panes == "A2"
            assert all(cell.font.name == "Arial" for row in sheet.iter_rows() for cell in row)
            assert not any(
                isinstance(cell.value, str) and cell.value.startswith("=")
                for row in sheet.iter_rows()
                for cell in row
            )


def test_generated_workbooks_support_api_import_and_planning(api_client: TestClient, tmp_path: Path) -> None:
    build_templates(tmp_path)
    scenario = "sim_oir_cuas"
    scenario_dir = tmp_path / scenario
    project_id = _create_project(api_client, "sim-import-roundtrip")

    task_result = _upload(api_client, project_id, "upload-task-units", scenario_dir / "task_units.xlsx")
    equipment_result = _upload(api_client, project_id, "upload-equipment-groups", scenario_dir / "equipment_groups.xlsx")
    rules_result = _upload(api_client, project_id, "upload-spectrum-rules", scenario_dir / "spectrum_rules.xlsx")
    assert task_result["count"] == 6
    assert equipment_result["count"] > 0
    assert rules_result["count"] > 0

    validation = api_client.post(f"/api/projects/{project_id}/validate-task")
    assert validation.status_code == 200, validation.text
    assert validation.json()["ok"] is True

    plan = api_client.post(
        f"/api/projects/{project_id}/task-plan",
        json={"objective": "task_assurance", "constraint_weights": {}, "strategy_profile": "balanced"},
    )
    assert plan.status_code == 200, plan.text
    assert plan.json()["status"] == "success"


def test_research_workbook_is_rejected_without_erasing_existing_rules(api_client: TestClient) -> None:
    project_id = _create_project(api_client, "reject-public-research-workbook")
    generated = api_client.post(f"/api/projects/{project_id}/generate-task-demo?scenario=sim_mosul_urban")
    assert generated.status_code == 200, generated.text

    before = api_client.get(f"/api/projects/{project_id}/task-data")
    assert before.status_code == 200
    before_rule_ids = [item["rule_id"] for item in before.json()["spectrum_rules"]]
    assert before_rule_ids

    research_workbook = ROOT / "docs" / "research" / "us_public_battle_spectrum_baseline.xlsx"
    with research_workbook.open("rb") as handle:
        rejected = api_client.post(
            f"/api/projects/{project_id}/upload-spectrum-rules",
            files={"file": (research_workbook.name, handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert rejected.status_code == 422
    assert "未覆盖现有项目数据" in rejected.text

    after = api_client.get(f"/api/projects/{project_id}/task-data")
    assert after.status_code == 200
    assert [item["rule_id"] for item in after.json()["spectrum_rules"]] == before_rule_ids


def test_invalid_full_schema_workbook_is_rejected_without_erasing_rules(
    api_client: TestClient,
    tmp_path: Path,
) -> None:
    build_templates(tmp_path)
    project_id = _create_project(api_client, "reject-invalid-full-schema")
    generated = api_client.post(f"/api/projects/{project_id}/generate-task-demo?scenario=sim_oir_cuas")
    assert generated.status_code == 200, generated.text
    before = api_client.get(f"/api/projects/{project_id}/task-data").json()["spectrum_rules"]

    invalid_path = tmp_path / "sim_oir_cuas" / "spectrum_rules.xlsx"
    workbook = load_workbook(invalid_path)
    sheet = workbook.active
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["end_mhz"], 0)
    workbook.save(invalid_path)

    with invalid_path.open("rb") as handle:
        rejected = api_client.post(
            f"/api/projects/{project_id}/upload-spectrum-rules",
            files={"file": (invalid_path.name, handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert rejected.status_code == 422
    assert "终止频率必须大于起始频率" in rejected.text
    after = api_client.get(f"/api/projects/{project_id}/task-data").json()["spectrum_rules"]
    assert [item["rule_id"] for item in after] == [item["rule_id"] for item in before]


def test_optional_columns_remain_backward_compatible(api_client: TestClient, tmp_path: Path) -> None:
    build_templates(tmp_path)
    project_id = _create_project(api_client, "optional-column-compatibility")
    generated = api_client.post(f"/api/projects/{project_id}/generate-task-demo?scenario=sim_oir_cuas")
    assert generated.status_code == 200, generated.text

    compatible_path = tmp_path / "sim_oir_cuas" / "spectrum_rules.xlsx"
    workbook = load_workbook(compatible_path)
    sheet = workbook.active
    headers = {cell.value: cell.column for cell in sheet[1]}
    for column_name in ("source", "reason"):
        sheet.delete_cols(headers[column_name])
        headers = {cell.value: cell.column for cell in sheet[1]}
    workbook.save(compatible_path)

    result = _upload(api_client, project_id, "upload-spectrum-rules", compatible_path)
    assert set(result["missing_columns"]) == {"source", "reason"}


def test_invalid_numeric_text_is_not_silently_defaulted(api_client: TestClient, tmp_path: Path) -> None:
    build_templates(tmp_path)
    project_id = _create_project(api_client, "reject-invalid-numeric-text")
    generated = api_client.post(f"/api/projects/{project_id}/generate-task-demo?scenario=sim_oir_cuas")
    assert generated.status_code == 200, generated.text
    before = api_client.get(f"/api/projects/{project_id}/task-data").json()["equipment_groups"]

    invalid_path = tmp_path / "sim_oir_cuas" / "equipment_groups.xlsx"
    workbook = load_workbook(invalid_path)
    sheet = workbook.active
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["count"], "not-a-number")
    workbook.save(invalid_path)

    with invalid_path.open("rb") as handle:
        rejected = api_client.post(
            f"/api/projects/{project_id}/upload-equipment-groups",
            files={"file": (invalid_path.name, handle, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert rejected.status_code == 422
    assert "数字格式无效" in rejected.text
    after = api_client.get(f"/api/projects/{project_id}/task-data").json()["equipment_groups"]
    assert [item["equipment_group_id"] for item in after] == [item["equipment_group_id"] for item in before]
