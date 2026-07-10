from __future__ import annotations

import hashlib
import json
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app.main import app
from backend.app.schemas import (
    EquipmentLibraryItemResponse,
    TaskObjectiveResponse,
    TaskSampleCatalogResponse,
    TaskScenarioResponse,
)
from backend.app.services import equipment_planning, task_catalog
from backend.app.services.excel_io import (
    EQUIPMENT_GROUP_COLUMNS,
    RULE_COLUMNS,
    SPECTRUM_RULE_COLUMNS,
    STATION_COLUMNS,
    TASK_UNIT_COLUMNS,
)


CATALOG_NAMES = (
    "TASK_OBJECTIVES",
    "TRIAL_METRIC_KEYS",
    "TRIAL_METRIC_LABELS",
    "EQUIPMENT_LIBRARY",
    "TASK_SCENARIOS",
    "SAMPLE_GENERATOR_PRESETS",
    "BATCH_PERFORMANCE_SCALES",
    "PLANNING_WEIGHT_TEMPLATES",
    "PARAMETRIC_SAMPLE_DEFAULTS",
)

CATALOG_SHA256 = "1fb8dd4dd1225fa429566ea69c4ac28c13f89e9d3937a61efc5b83ba7a90d8ae"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_catalog_content_and_legacy_imports_are_stable() -> None:
    catalog = {name: getattr(task_catalog, name) for name in CATALOG_NAMES}
    normalized = json.dumps(catalog, ensure_ascii=False, sort_keys=True, separators=(",", ":"))

    assert hashlib.sha256(normalized.encode("utf-8")).hexdigest() == CATALOG_SHA256
    for name, value in catalog.items():
        assert getattr(equipment_planning, name) == value


def test_catalog_endpoints_keep_their_public_contract() -> None:
    client = TestClient(app)

    objectives = client.get("/api/task-objectives")
    equipment = client.get("/api/equipment-library")
    scenarios = client.get("/api/task-scenarios")
    samples = client.get("/api/task-sample-catalog")

    for response in (objectives, equipment, scenarios, samples):
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("application/json")

    assert len(objectives.json()) == 10
    assert len(equipment.json()) == 17
    assert len(scenarios.json()) == 12
    assert objectives.json() == task_catalog.TASK_OBJECTIVES
    assert equipment.json() == task_catalog.EQUIPMENT_LIBRARY
    assert scenarios.json() == task_catalog.TASK_SCENARIOS
    assert samples.json() == equipment_planning.task_sample_catalog()
    assert set(samples.json()) == {
        "presets",
        "task_unit_profiles",
        "equipment_types",
        "batch_scales",
        "weight_templates",
        "parametric_defaults",
        "objective_count",
    }
    assert any(item["key"] == "balanced" for item in samples.json()["weight_templates"])


def test_catalog_routes_declare_structured_response_models() -> None:
    routes = {route.path: route for route in app.routes}

    assert routes["/api/task-objectives"].response_model == list[TaskObjectiveResponse]
    assert routes["/api/equipment-library"].response_model == list[EquipmentLibraryItemResponse]
    assert routes["/api/task-scenarios"].response_model == list[TaskScenarioResponse]
    assert routes["/api/task-sample-catalog"].response_model is TaskSampleCatalogResponse


def test_catalog_openapi_uses_structured_component_references() -> None:
    schema = TestClient(app).get("/openapi.json").json()
    paths = schema["paths"]

    list_components = {
        "/api/task-objectives": "TaskObjectiveResponse",
        "/api/equipment-library": "EquipmentLibraryItemResponse",
        "/api/task-scenarios": "TaskScenarioResponse",
    }
    for path, component in list_components.items():
        response_schema = paths[path]["get"]["responses"]["200"]["content"]["application/json"]["schema"]
        assert response_schema["type"] == "array"
        assert response_schema["items"]["$ref"] == f"#/components/schemas/{component}"

    sample_schema = paths["/api/task-sample-catalog"]["get"]["responses"]["200"]["content"]["application/json"]["schema"]
    assert sample_schema["$ref"] == "#/components/schemas/TaskSampleCatalogResponse"

    components = schema["components"]["schemas"]
    assert components["TaskSampleCatalogResponse"]["properties"]["task_unit_profiles"]["items"]["$ref"] == (
        "#/components/schemas/TaskUnitProfileResponse"
    )
    assert components["TaskSampleCatalogResponse"]["properties"]["weight_templates"]["items"]["$ref"] == (
        "#/components/schemas/PlanningWeightTemplateResponse"
    )
    assert components["PlanningWeightTemplateResponse"]["properties"]["weights"]["$ref"] == (
        "#/components/schemas/PlanningWeightsResponse"
    )


@pytest.mark.parametrize(
    ("path", "filename", "columns"),
    [
        ("/api/templates/stations.xlsx", "stations_template.xlsx", STATION_COLUMNS),
        ("/api/templates/rules.xlsx", "rules_template.xlsx", RULE_COLUMNS),
        ("/api/templates/task-units.xlsx", "task_units_template.xlsx", TASK_UNIT_COLUMNS),
        ("/api/templates/equipment-groups.xlsx", "equipment_groups_template.xlsx", EQUIPMENT_GROUP_COLUMNS),
        ("/api/templates/spectrum-rules.xlsx", "spectrum_rules_template.xlsx", SPECTRUM_RULE_COLUMNS),
    ],
)
def test_excel_template_contract(path: str, filename: str, columns: list[str]) -> None:
    response = TestClient(app).get(path)

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(XLSX_MEDIA_TYPE)
    assert response.headers["content-disposition"] == f'attachment; filename="{filename}"'
    assert response.content.startswith(b"PK")

    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert workbook.sheetnames == ["Sheet1"]
    assert list(next(workbook.active.iter_rows(values_only=True))) == columns


def test_catalog_routes_remain_in_openapi_and_allow_the_main_frontend_origin() -> None:
    client = TestClient(app)
    schema = client.get("/openapi.json").json()
    paths = schema["paths"]

    for path in (
        "/api/task-objectives",
        "/api/equipment-library",
        "/api/task-scenarios",
        "/api/task-sample-catalog",
    ):
        assert "get" in paths[path]

    preflight = client.options(
        "/api/task-objectives",
        headers={
            "Origin": "http://127.0.0.1:5173",
            "Access-Control-Request-Method": "GET",
        },
    )
    assert preflight.status_code == 200
    assert preflight.headers["access-control-allow-origin"] == "http://127.0.0.1:5173"
